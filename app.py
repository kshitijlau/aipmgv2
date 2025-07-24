# app.py (version 12.2 - Azure DevOps Formatting)

import streamlit as st
import google.generativeai as genai
from datetime import datetime
import docx
import io
import json
import re 
import traceback

# --- App Constants ---
MAX_FILE_SIZE_MB = 15
MAX_FILE_SIZE_BYTES = MAX_FILE_SIZE_MB * 1024 * 1024
MODEL_NAME = "gemini-2.5-pro" 

# --- Page Configuration ---
st.set_page_config(
    page_title="AI Requirements Assistant (Gemini)",
    page_icon="✨",
    layout="wide"
)

# --- Prompts ---

# --- TRANSCRIPT ANALYSIS PROMPTS (MULTI-PHASE) ---

# Prompt 1A: The Signals Intelligence Chunk Analysis Prompt for transcripts
CHUNK_ANALYSIS_PROMPT = """
**Persona:**
You are an AI Signals Intelligence (SIGINT) Analyst. Your sole mission is to meticulously analyze a small, decontextualized snippet of a longer conversation and extract every potential data point without judgment or synthesis. You are a specialist in identifying and categorizing raw information for later analysis by a different system. You are incapable of missing a requirement.

**Primary Directive:**
You will use a Chain-of-Thought process. For each sentence in the provided transcript chunk, you will first think about what signals are present, and only then will you add them to a final JSON object. Your ONLY output must be a single, valid JSON object. Do not output your thoughts, only the final JSON.

**Core Principles of Extraction:**
1.  **Deconstruct Every Sentence:** Do not read paragraphs as a whole. Analyze one sentence at a time. A single sentence can contain multiple signals. You must extract all of them.
2.  **Extract, Do Not Interpret:** Do not summarize or rephrase. Capture the essence of the statement, using direct quotes where possible.
3.  **No Signal Left Behind:** Your primary goal is to capture everything that *might* be relevant.
4.  **Attribute to Speaker:** If the text indicates who is speaking (e.g., "Sarah:", "[John]"), you MUST populate the `speaker` field. If it's unclear, use "Unknown".

**Signal Categorization Protocol:**
Categorize every extracted point into ONE of these: `Explicit_Requirement`, `Implicit_Requirement`, `Technical_Specification`, `UI_UX_Detail`, `Decision_Made`, `Action_Item`, `User_Pain_Point`, `Business_Goal`, `Open_Question`, `Identified_Risk`.

**Required JSON Output Format:**
```json
{{
  "extracted_signals": [
    {{
      "category": "ENUM(One of the categories from the protocol above)",
      "speaker": "STRING",
      "content": "STRING(The extracted statement or key phrase)",
      "priority_signal": "ENUM('High', 'Medium', 'Low')"
    }}
  ]
}}
```

Now, apply this rigorous process to the following transcript chunk.

---TRANSCRIPT CHUNK---
{chunk_text}
---TRANSCRIPT CHUNK---
"""

# Prompt 1B: The Master Architect Synthesis Prompt for transcripts
FINAL_SYNTHESIS_PROMPT = """
**Persona:**
You are an AI Master Architect. Your exclusive function is to receive structured, raw intelligence data (in JSON format) and synthesize it into a single, comprehensive, and strategically coherent Product Requirements Document (PRD) formatted for optimal display in systems like Azure DevOps wikis.

**Primary Directive:**
Your sole input is a list of JSON objects, where each object is an "extracted signal". Your only output must be a single, human-readable PRD formatted in Markdown. You must process **every single signal** from the input and ensure it is appropriately represented in the final document, using markdown tables for clarity where specified in the template.

**Core Principles of Synthesis:**
1.  **Input is Ground Truth:** The provided JSON data is your only source of information. Do not add requirements not supported by a signal.
2.  **De-duplicate & Synthesize:** Intelligently merge related signals into a single, cohesive requirement entry in the final PRD.
3.  **Structure is Paramount:** Follow the PRD template below with absolute precision. Use markdown tables exactly as shown.
4.  **Follow the Mapping Guide:** Use the protocol to determine where each signal category should be placed in the final document.

**Signal-to-PRD Mapping Protocol:**
-   **`Explicit_Requirement` / `Implicit_Requirement`:** These form the "Requirement" sections in the Work Breakdown.
-   **`UI_UX_Detail`:** Translate into FE Tasks or Acceptance Criteria.
-   **`Technical_Specification`:** Translate into BE Tasks, Acceptance Criteria, or NFRs.
-   **`User_Pain_Point` / `Business_Goal`:** Use for the Strategic Overview and the "so that I can..." part of user stories.
-   **`Decision_Made`:** Use to state definitive behavior in Acceptance Criteria.
-   **`Action_Item` / `Open_Question`:** Convert into line items in the "Open Questions & Action Items" section.
-   **`Identified_Risk`:** Add to the "Potential Risks" table for the relevant requirement.

**--- PRD TEMPLATE FOR AZURE DEVOPS WIKI ---**
*Generate the final document using this exact structure.*

# PRD: Lighthouse Platform - [Feature Name]
---
## 1. Strategic Overview
- **Feature Name:** [Synthesize a clear name from the signals]
- **User "Job to Be Done" (JTBD):** [Synthesize from `User_Pain_Point` and `Implicit_Requirement` signals: "When I [context], I want to [motivation], so I can [expected outcome]."]
- **Business Goal:** [Synthesize from `Business_Goal` signals.]
- **Success Metrics:** [Infer potential KPIs from business goals.]
---
## 2. Open Questions & Action Items
*Synthesized directly from `Open_Question` and `Action_Item` signals.*
- **[ ] Open Question:** [Content of `Open_Question` signal] - **Owner:** [Suggest a role]
- **[ ] Action Item:** [Content of `Action_Item` signal] - **Owner:** [Suggest a role]
---
## 3. Non-Functional Requirements (NFRs)
*Synthesize from any relevant `Technical_Specification` signals that are global in nature. Format as a markdown table.*

| Category      | Requirement                                                | Metric/Standard                  |
|---------------|------------------------------------------------------------|----------------------------------|
| Performance   | [e.g., API Response Time]                                  | [e.g., 95% of responses < 500ms] |
| Security      | [e.g., Authentication]                                     | [e.g., All endpoints are secured]|
| Accessibility | [e.g., Keyboard Navigation]                                | [e.g., WCAG 2.1 AA Compliant]    |

---
## 4. Epic & Work Breakdown Structure
*A complete deconstruction of the work required, built by mapping all relevant signals.*
### ### Epic: [Synthesize a high-level Epic title]
*This epic covers all work required for the discussed feature set.*
---
### Requirement: [Title synthesized from one or more `Explicit_Requirement` signals]
- **User Story:** [Synthesize from signals]
- **Priority:** [Determine from `priority_signal` values]
- **Acceptance Criteria:**
    - [ ] [Synthesize from `UI_UX_Detail`, `Technical_Specification`, and `Decision_Made` signals.]
    - [ ] [Add another criterion...]

**Implementation Tasks:**
*Format as a markdown table.*
| Discipline | Task Description                                       | Notes                  |
|------------|--------------------------------------------------------|------------------------|
| Frontend   | [Synthesize from one or more `UI_UX_Detail` signals.]  | [Any additional notes] |
| Backend    | [Synthesize from one or more `Technical_Specification` signals.] | [Any additional notes] |
| QA         | [Create a specific task to verify the acceptance criteria.] | [e.g., "End-to-end test"] |

**Potential Risks:**
*Format as a markdown table.*
| Risk Category | Description                                         | Mitigation Strategy         |
|---------------|-----------------------------------------------------|-----------------------------|
| [e.g., Technical] | [Synthesize from any `Identified_Risk` signals.] | [Suggest a mitigation plan] |

---
*(Continue this structure for every other requirement identified in the input signals...)*
"""


# --- PM NOTES ANALYSIS PROMPT (SINGLE-PHASE) ---
PM_NOTES_PROMPT = """
**Persona:**
You are an AI Product Strategist. Your expertise is in taking a product manager's rough, unstructured notes and transforming them into a comprehensive, engineering-ready Product Requirements Document (PRD), formatted for optimal display in systems like Azure DevOps wikis.

**Primary Directive:**
Analyze the provided PM notes. Your goal is to first deconstruct the information and rebuild it into a structured PRD using markdown tables for clarity. After structuring the notes, you will then provide a section with strategic suggestions for improving the feature.

**Core Principles of Interpretation & Analysis:**
1.  **Structure from Chaos:** Your first job is to impose the PRD structure onto the notes. Group related points into a single requirement.
2.  **Identify the Gaps:** If the notes are unclear or missing key details, you must create logical, inferred placeholders and flag them in the "Open Questions" section.
3.  **Be Exhaustive:** Do not discard any point from the notes. Every idea or feature mentioned must be translated into a corresponding section in the PRD.
4.  **Think Strategically:** In the final "Suggestions" section, think beyond the notes. Consider the user experience, potential edge cases, and future scalability. Provide concrete ideas to make the feature truly excellent.

**--- PRD TEMPLATE FOR AZURE DEVOPS WIKI ---**
*You must generate the final document using this exact structure.*

# PRD: Lighthouse Platform - [Feature Name]
---
## 1. Strategic Overview
- **Feature Name:** [Determine a clear name from the notes]
- **User "Job to Be Done" (JTBD):** [Determine from the notes: "When I [context], I want to [motivation], so I can [expected outcome]."]
- **Business Goal:** [Determine from the notes.]
- **Success Metrics:** [Infer potential KPIs from the goals.]
---
## 2. Open Questions & Action Items
*List all points that need clarification based on your analysis of the notes.*
- **[ ] Open Question:** [e.g., "What formats should be supported for export (CSV, PDF, etc.)?"] - **Owner:** @Product
- **[ ] Action Item:** [e.g., "Confirm performance requirements for large data exports."] - **Owner:** @Engineering
---
## 3. Non-Functional Requirements (NFRs)
*Infer any NFRs mentioned or implied in the notes. Format as a markdown table.*

| Category      | Requirement                                                | Metric/Standard                  |
|---------------|------------------------------------------------------------|----------------------------------|
| Performance   | [e.g., API Response Time]                                  | [e.g., Notes mention "must be fast"] |
| Security      | [e.g., Authentication]                                     | [e.g., Notes imply "only authenticated users"]|
| Accessibility | [e.g., Keyboard Navigation]                                | [e.g., Inferred: WCAG 2.1 AA Compliant]    |

---
## 4. Epic & Work Breakdown Structure
*A complete deconstruction of the work required, built from the notes.*
### ### Epic: [Create a high-level Epic title from the notes]
*This epic covers all work for the features described.*
---
### Requirement: [Title for the first requirement identified in the notes]
- **User Story:** [Write a full user story based on the note]
- **Priority:** [Assign a logical priority, e.g., P1-High]
- **Acceptance Criteria:**
    - [ ] [Create logical acceptance criteria for the requirement.]
    - [ ] [Add another criterion...]

**Implementation Tasks:**
*Format as a markdown table.*
| Discipline | Task Description                                       | Notes                  |
|------------|--------------------------------------------------------|------------------------|
| Frontend   | [Create a specific frontend task.]                     | [Any additional notes] |
| Backend    | [Create a specific backend task.]                      | [Any additional notes] |
| QA         | [Create a specific testing task.]                      | [e.g., "End-to-end test"] |

**Potential Risks:**
*Format as a markdown table.*
| Risk Category | Description                                         | Mitigation Strategy         |
|---------------|-----------------------------------------------------|-----------------------------|
| [e.g., Technical] | [Identify any potential risks based on the notes.] | [Suggest a mitigation plan] |

---
*(Continue this structure for every other feature or requirement identified in the notes...)*

---
## 5. Strategic Suggestions & Future Enhancements
*Your analysis and ideas for making this feature better.*
- **Immediate Improvements (V1.0):**
    - **Suggestion:** [Provide a specific, actionable idea to improve the initial version of the feature. e.g., "Add a 'Recent Exports' link in the modal so users can quickly re-download files."]
    - **Rationale:** [Explain why this suggestion adds value.]
- **Future Roadmap Ideas (V2.0 and beyond):**
    - **Suggestion:** [Provide a bigger-picture idea for the future. e.g., "Implement scheduled, recurring exports that can be emailed to users automatically."]
    - **Rationale:** [Explain how this enhancement addresses a deeper user need or business goal.]
"""


# --- Functions ---

def get_text_from_docx(docx_bytes):
    """Extracts text from a .docx file."""
    doc = docx.Document(io.BytesIO(docx_bytes))
    full_text = []
    for para in doc.paragraphs:
        full_text.append(para.text)
    return '\n'.join(full_text)

def get_text_chunks(text, chunk_size=12000, overlap=500):
    """Splits text into overlapping chunks. Increased chunk size for Gemini."""
    if len(text) <= chunk_size:
        return [text]
    
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunks.append(text[start:end])
        start += chunk_size - overlap
    return chunks

def get_gemini_response(model, prompt_parts, is_json_output=False):
    """Generic function to get a response from the Gemini model."""
    config = {"temperature": 0.0}
    if is_json_output:
        config["response_mime_type"] = "application/json"
    
    response = model.generate_content(prompt_parts, generation_config=config)
    return response.text

def process_long_transcript(model, transcript_text):
    """Orchestrates the chunking and synthesis process using Gemini for transcripts."""
    chunks = get_text_chunks(transcript_text)
    all_signals = []
    progress_bar = st.progress(0, text="Phase 1: Analyzing transcript chunks...")
    total_chunks = len(chunks)

    for i, chunk in enumerate(chunks):
        prompt_for_chunk = CHUNK_ANALYSIS_PROMPT.format(chunk_text=chunk)
        try:
            raw_response = get_gemini_response(model, prompt_for_chunk, is_json_output=True)
            data = json.loads(raw_response)
            if isinstance(data, dict) and "extracted_signals" in data and isinstance(data["extracted_signals"], list):
                all_signals.extend(data["extracted_signals"])
            elif isinstance(data, list):
                all_signals.extend(data)
            else:
                st.warning(f"JSON from chunk {i+1} has an unexpected structure. Skipping.")
        except Exception as e:
            st.warning(f"Could not process chunk {i+1} due to an error: {e}. Skipping.")
        progress_bar.progress((i + 1) / total_chunks, text=f"Phase 1: Analyzed chunk {i+1} of {total_chunks}")
        
    progress_bar.progress(1.0, text="Phase 2: Synthesizing final document from all signals...")
    
    if not all_signals:
        st.error("Analysis complete, but no valid requirements could be extracted. The final PRD cannot be generated.", icon="🚨")
        return None

    combined_signals_json = json.dumps({"all_extracted_signals": all_signals}, indent=2)
    final_prd = get_gemini_response(model, [FINAL_SYNTHESIS_PROMPT, combined_signals_json])
    
    progress_bar.empty()
    return final_prd

def process_pm_notes(model, notes_text):
    """Processes PM notes directly into a PRD in a single pass."""
    return get_gemini_response(model, [PM_NOTES_PROMPT, notes_text])


# --- Main Application ---
st.title("✨ AI Requirements Assistant (Gemini Pro)")
st.markdown(f"Upload a document (`.txt` or `.docx`, max {MAX_FILE_SIZE_MB}MB) to generate a comprehensive requirements document.")

# --- Secrets Check and Client Initialization ---
try:
    GEMINI_API_KEY = st.secrets["google_generativeai"]["api_key"]
    if not GEMINI_API_KEY:
        st.error("Google Gemini API key is missing. Please add it to your Streamlit secrets.", icon="🚨")
        st.stop()
    
    genai.configure(api_key=GEMINI_API_KEY)
    gemini_model = genai.GenerativeModel(MODEL_NAME)

except KeyError:
    st.error("Google Gemini credentials are not set correctly. Make sure you have a [google_generativeai] section with an 'api_key' in your secrets.", icon="🚨")
    st.stop()
except Exception as e:
    st.error(f"Failed to initialize Google Gemini client: {e}", icon="🚨")
    st.stop()

# --- User Interface ---
st.subheader("1. Select Document Type")
input_type = st.radio(
    "What type of document are you uploading?",
    ("Meeting Transcript", "Product Manager's Notes"),
    horizontal=True,
    label_visibility="collapsed"
)

st.subheader("2. Upload Your Document")
uploaded_file = st.file_uploader(
    "Upload file",
    type=['txt', 'docx'],
    label_visibility="collapsed"
)

if uploaded_file:
    if uploaded_file.size > MAX_FILE_SIZE_BYTES:
        st.error(f"File size exceeds the {MAX_FILE_SIZE_MB}MB limit. Please upload a smaller file.", icon="🚨")
    else:
        # Read file content
        raw_text = ""
        if uploaded_file.type == "text/plain":
            raw_text = uploaded_file.getvalue().decode("utf-8")
        elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            raw_text = get_text_from_docx(uploaded_file.getvalue())
        
        st.subheader("3. Generate Document")
        if st.button(f"🚀 Generate PRD from {input_type}"):
            if not raw_text.strip():
                st.warning("The uploaded document appears to be empty. Please upload a file with content.", icon="⚠️")
            else:
                analysis_result = None
                try:
                    if input_type == "Meeting Transcript":
                        with st.spinner("Starting multi-phase analysis of transcript... This may take several minutes."):
                            analysis_result = process_long_transcript(
                                model=gemini_model,
                                transcript_text=raw_text
                            )
                    elif input_type == "Product Manager's Notes":
                        with st.spinner("Analyzing PM notes and generating PRD..."):
                           analysis_result = process_pm_notes(
                               model=gemini_model,
                               notes_text=raw_text
                           )
                    
                    if analysis_result:
                        st.success("Analysis Complete!", icon="🎉")

                        with st.expander("View Full Requirements Document", expanded=True):
                            st.markdown(analysis_result)
                        
                        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                        file_name = f"Lighthouse_Requirements_{timestamp}.md"
                        
                        st.download_button(
                            label="⬇️ Download Requirements as .md File",
                            data=analysis_result.encode('utf-8'),
                            file_name=file_name,
                            mime='text/markdown'
                        )
                except Exception as e:
                    st.error(f"An error occurred during processing: {e}", icon="🚨")
                    traceback.print_exc()
